"""Find, download, and parse the World Bank monthly commodity price workbook.

The World Bank publishes CMO-Historical-Data-Monthly.xlsx (the "Pink Sheet"
monthly history) on its Commodity Markets page. The file URL changes with each
release, so the collector reads the page and finds the link.

Layout of the "Monthly Prices" sheet, as documented by the World Bank and seen in
past releases:

    a few title rows ("World Bank Commodity Price Data (The Pink Sheet)",
                      "Updated on October 02, 2026", ...)
    a row of series names     ("Crude oil, average", "Crude oil, Brent", ...)
    a row of units            ("($/bbl)", "($/mmbtu)", ...)
    (often) a row of codes    ("CRUDE_PETRO", "CRUDE_BRENT", ...)
    data rows                 ("1960M01", 1.63, 1.63, ...)

The parser does not hard-code row numbers. It finds the first "YYYYMmm" row and
classifies the rows above it. Missing values appear as "..", "…", blanks, or text,
and become None.

Only runtime dependency: openpyxl.
"""

from __future__ import annotations

import datetime as _dt
import hashlib
import io
import json
import pathlib
import re
import urllib.parse
import urllib.request

from .analytics import classify

LICENSE = {
    "name": "Creative Commons Attribution 4.0 International (CC BY 4.0)",
    "url": "https://datacatalog.worldbank.org/public-licenses?fragment=cc",
    "dataset_url": "https://datacatalog.worldbank.org/search/dataset/0038238/commodity-prices-history-and-projections",
    "attribution": (
        "Source: World Bank, Commodity Price Data (The Pink Sheet), "
        "https://www.worldbank.org/en/research/commodity-markets. "
        "Licensed under CC BY 4.0. Reformatted and derived statistics computed by Commodity Pulse; "
        "the World Bank does not endorse this tool."
    ),
}

SOURCE_PAGE = "https://www.worldbank.org/en/research/commodity-markets"
ALLOWED_HOSTS = ("worldbank.org", "thedocs.worldbank.org", "www.worldbank.org", "pubdocs.worldbank.org")
MAX_PAGE_BYTES = 5 * 1024 * 1024
MAX_XLSX_BYTES = 60 * 1024 * 1024
SHEET_NAME = "Monthly Prices"
USER_AGENT = "commodity-pulse (+https://github.com/)"

DATE_RE = re.compile(r"^\s*(\d{4})\s*M\s*(\d{1,2})\s*$", re.I)
CODE_RE = re.compile(r"^[A-Z][A-Z0-9_]*$")
HREF_RE = re.compile(r"""href\s*=\s*["']([^"']*CMO-Historical-Data-Monthly\.xlsx[^"']*)["']""", re.I)
UPDATED_RE = re.compile(r"updated\s+on\s+([A-Za-z]+\s+\d{1,2},?\s+\d{4})", re.I)

# Fallback codes when a workbook has no code row. Keys are normalized names.
KNOWN_NAMES = {
    "crude oil average": "CRUDE_PETRO", "crude oil brent": "CRUDE_BRENT", "crude oil dubai": "CRUDE_DUBAI",
    "crude oil wti": "CRUDE_WTI", "coal australian": "COAL_AUS", "coal australia": "COAL_AUS",
    "coal south african": "COAL_SAFRICA", "coal south africa": "COAL_SAFRICA",
    "natural gas us": "NGAS_US", "natural gas europe": "NGAS_EUR",
    "liquefied natural gas japan": "NGAS_JP", "natural gas index": "INATGAS",
    "urea": "UREA_EE_BULK", "urea e europe": "UREA_EE_BULK", "dap": "DAP", "tsp": "TSP",
    "phosphate rock": "PHOSROCK", "potassium chloride": "POTASH",
    "aluminum": "ALUMINUM", "copper": "COPPER", "iron ore cfr spot": "IRON_ORE", "iron ore": "IRON_ORE",
    "lead": "LEAD", "nickel": "NICKEL", "tin": "TIN", "zinc": "ZINC",
    "gold": "GOLD", "silver": "SILVER", "platinum": "PLATINUM",
    "maize": "MAIZE", "soybeans": "SOYBEANS", "wheat us hrw": "WHEAT_US_HRW", "wheat us srw": "WHEAT_US_SRW",
    "cocoa": "COCOA", "coffee arabica": "COFFEE_ARABIC", "coffee robusta": "COFFEE_ROBUS",
    "sugar world": "SUGAR_WLD", "cotton a index": "COTTON_A_INDX", "palm oil": "PALM_OIL",
}


class CollectError(RuntimeError):
    """Collector failure with a stable ``code``."""

    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code


# ---------------------------------------------------------------- locate + download

def find_workbook_url(html: str, base: str = SOURCE_PAGE) -> str:
    """Return the absolute URL of CMO-Historical-Data-Monthly.xlsx found in the page HTML."""
    for m in HREF_RE.finditer(html or ""):
        url = urllib.parse.urljoin(base, m.group(1).replace("&amp;", "&").strip())
        check_url(url)
        return url
    raise CollectError("no_link", "No CMO-Historical-Data-Monthly.xlsx link found on " + base + ".")


def check_url(url: str) -> None:
    p = urllib.parse.urlparse(url)
    host = (p.hostname or "").lower()
    if p.scheme != "https":
        raise CollectError("bad_url", "Refusing non-https URL: " + url)
    if not any(host == h or host.endswith("." + h) for h in ALLOWED_HOSTS):
        raise CollectError("bad_url", "Refusing URL outside worldbank.org: " + url)


class _CheckedRedirects(urllib.request.HTTPRedirectHandler):
    """Refuse a redirect before following it unless the target passes check_url."""

    def redirect_request(self, req, fp, code, msg, headers, newurl):
        check_url(newurl)
        return super().redirect_request(req, fp, code, msg, headers, newurl)


_OPENER = urllib.request.build_opener(_CheckedRedirects)

# Series every real Pink Sheet monthly workbook carries. If any is missing after
# parsing, the layout changed and the collector stops instead of writing bad data.
REQUIRED_SERIES = ("CRUDE_BRENT", "CRUDE_WTI", "NGAS_US", "NGAS_EUR", "GOLD", "SILVER", "MAIZE", "UREA_EE_BULK")


def _get(url: str, limit: int, timeout: float = 60.0) -> bytes:
    check_url(url)
    req = urllib.request.Request(url, headers={"User-Agent": USER_AGENT})
    with _OPENER.open(req, timeout=timeout) as r:  # noqa: S310 (https + host allowlist above)
        final = r.geturl()
        check_url(final)
        data = r.read(limit + 1)
    if len(data) > limit:
        raise CollectError("too_big", "Download exceeds " + str(limit) + " bytes: " + url)
    return data


def download(page_url: str = SOURCE_PAGE) -> tuple[str, bytes]:
    html = _get(page_url, MAX_PAGE_BYTES).decode("utf-8", "replace")
    url = find_workbook_url(html, page_url)
    return url, _get(url, MAX_XLSX_BYTES)


# ---------------------------------------------------------------- parse

def _text(v) -> str:
    if v is None:
        return ""
    return str(v).strip()


def _norm_name(s: str) -> str:
    s = s.lower().replace("**", " ")
    s = re.sub(r"\(.*?\)", " ", s)
    return " ".join(re.findall(r"[a-z0-9]+", s))


def _slug(s: str) -> str:
    return "_".join(re.findall(r"[A-Z0-9]+", s.upper())) or "SERIES"


def clean_unit(s: str) -> str:
    s = _text(s)
    if s.startswith("(") and s.endswith(")"):
        s = s[1:-1].strip()
    return s


def parse_number(v):
    """Cell value to float or None. Accepts numbers and numeric strings."""
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        f = float(v)
        return f if f == f and f not in (float("inf"), float("-inf")) else None
    s = str(v).strip().replace(",", "")
    if not s:
        return None
    try:
        f = float(s)
    except ValueError:
        return None
    return f if f == f and f not in (float("inf"), float("-inf")) else None


def parse_date_cell(v):
    """'1960M01' to '1960-01'. Also accepts datetime cells. None if not a month."""
    if isinstance(v, (_dt.datetime, _dt.date)):
        return f"{v.year:04d}-{v.month:02d}"
    m = DATE_RE.match(_text(v))
    if not m:
        return None
    mo = int(m.group(2))
    if mo < 1 or mo > 12:
        return None
    return f"{int(m.group(1)):04d}-{mo:02d}"


def parse_release(text: str):
    m = UPDATED_RE.search(text or "")
    if not m:
        return None
    raw = m.group(1).replace(",", "")
    raw = re.sub(r"\s+", " ", raw)
    for fmt in ("%B %d %Y", "%b %d %Y"):
        try:
            return _dt.datetime.strptime(raw, fmt).date().isoformat()
        except ValueError:
            pass
    return None


def parse_rows(rows: list[list]) -> dict:
    """Parse the sheet given as a list of row lists (cell values).

    Returns {"release": ISO date or None, "rows": tidy rows}. Tidy row keys:
    series_id, name, unit, date (YYYY-MM), value (float). Missing cells are dropped.
    """
    first_data = None
    for i, r in enumerate(rows):
        if r and parse_date_cell(r[0]) is not None:
            first_data = i
            break
    if first_data is None:
        raise CollectError("no_data", "No data rows like 1960M01 found in column A.")
    header = rows[:first_data]
    width = max((len(r) for r in rows), default=0)

    release = None
    for r in header:
        for c in r:
            release = release or parse_release(_text(c))

    def cells(r):
        return [_text(r[j]) if j < len(r) else "" for j in range(1, width)]

    code_row = unit_row = name_row = None
    for i in range(len(header) - 1, -1, -1):
        cs = [c for c in cells(header[i]) if c]
        if len(cs) < 2 and width > 2:
            continue
        if not cs:
            continue
        if code_row is None and sum(1 for c in cs if CODE_RE.match(c)) >= 0.6 * len(cs) and unit_row is None:
            code_row = i
            continue
        if unit_row is None and sum(1 for c in cs if c.startswith("(") or "$" in c or "/" in c) >= 0.6 * len(cs):
            unit_row = i
            continue
        if name_row is None and (unit_row is not None or code_row is not None):
            name_row = i
            break
    if name_row is None:
        # No units/code rows recognized: take the nearest text row above the data.
        for i in range(len(header) - 1, -1, -1):
            if sum(1 for c in cells(header[i]) if c) >= 1:
                name_row = i
                break
    if name_row is None:
        raise CollectError("no_header", "No series-name row found above the data.")

    names = cells(header[name_row])
    units = cells(header[unit_row]) if unit_row is not None else [""] * len(names)
    codes = cells(header[code_row]) if code_row is not None else [""] * len(names)

    columns = []
    used = set()
    for j, name in enumerate(names):
        name = re.sub(r"\s+", " ", name.replace("**", "")).strip()
        if not name:
            continue
        code = codes[j] if j < len(codes) and codes[j] else KNOWN_NAMES.get(_norm_name(name), _slug(name))
        base, k = code, 2
        while code in used:
            code = f"{base}_{k}"
            k += 1
        used.add(code)
        columns.append((j + 1, code, name, clean_unit(units[j] if j < len(units) else "")))
    if not columns:
        raise CollectError("no_header", "The series-name row is empty.")

    tidy = []
    seen_dates = set()
    for r in rows[first_data:]:
        if not r:
            continue
        d = parse_date_cell(r[0])
        if d is None:
            continue
        if d in seen_dates:
            raise CollectError("dup_date", "Month " + d + " appears twice.")
        seen_dates.add(d)
        for col, code, name, unit in columns:
            v = parse_number(r[col]) if col < len(r) else None
            if v is None:
                continue
            tidy.append({"series_id": code, "name": name, "unit": unit, "date": d, "value": v})
    if not tidy:
        raise CollectError("no_data", "Data rows contain no numeric values.")
    return {"release": release, "rows": tidy}


def parse_workbook(data: bytes, sheet: str = SHEET_NAME) -> dict:
    """Parse workbook bytes. Returns parse_rows output."""
    import openpyxl  # runtime dependency, imported here so analytics works without it

    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:  # openpyxl raises several types for a bad file
        raise CollectError("bad_workbook", "Not a readable .xlsx workbook: " + type(e).__name__) from None
    try:
        name = sheet if sheet in wb.sheetnames else next(
            (s for s in wb.sheetnames if s.strip().lower() == sheet.lower()), None)
        if name is None:
            raise CollectError("no_sheet", "Sheet '" + sheet + "' not found. Sheets: " + ", ".join(wb.sheetnames))
        rows = [list(r) for r in wb[name].iter_rows(values_only=True)]
    finally:
        wb.close()
    return parse_rows(rows)


# ---------------------------------------------------------------- tidy -> compact

def _month_index(d: str) -> int:
    return int(d[:4]) * 12 + int(d[5:7]) - 1


def _month_str(i: int) -> str:
    return f"{i // 12:04d}-{i % 12 + 1:02d}"


def to_dataset(tidy: list[dict]) -> dict:
    """Tidy rows to the compact per-series form used by data/monthly.json.

    Each series: id, name, unit, group, start (YYYY-MM), values (one per month,
    None where missing). Series keep first-seen order.
    """
    order = []
    by_id = {}
    for r in tidy:
        sid = r["series_id"]
        if sid not in by_id:
            by_id[sid] = {"name": r["name"], "unit": r["unit"], "points": {}}
            order.append(sid)
        by_id[sid]["points"][_month_index(r["date"])] = float(r["value"])
    series = []
    for sid in order:
        pts = by_id[sid]["points"]
        lo, hi = min(pts), max(pts)
        series.append({
            "id": sid,
            "name": by_id[sid]["name"],
            "unit": by_id[sid]["unit"],
            "group": classify(sid, by_id[sid]["name"]),
            "start": _month_str(lo),
            "values": [pts.get(i) for i in range(lo, hi + 1)],
        })
    return {"series": series}


def from_dataset(ds: dict) -> list[dict]:
    """Compact dataset back to tidy rows (missing months dropped)."""
    out = []
    for s in ds.get("series", []):
        base = _month_index(s["start"])
        for k, v in enumerate(s["values"]):
            if v is None:
                continue
            out.append({"series_id": s["id"], "name": s.get("name", s["id"]), "unit": s.get("unit", ""),
                        "date": _month_str(base + k), "value": v})
    return out


def dumps_compact(obj) -> str:
    return json.dumps(obj, separators=(",", ":"), ensure_ascii=False, allow_nan=False)


def data_hash(ds: dict) -> str:
    return hashlib.sha256(dumps_compact(ds).encode("utf-8")).hexdigest()


# ---------------------------------------------------------------- write

def write_outputs(ds: dict, out_dir: pathlib.Path, source_url: str, release, fetched_at: str | None = None,
                  extra_meta: dict | None = None) -> bool:
    """Write monthly.json and meta.json if the data changed. Returns True if written.

    "Changed" means a different release date or different data. A re-fetch of the
    same release leaves both files untouched, so the workflow makes no commit.
    """
    out_dir = pathlib.Path(out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    h = data_hash(ds)
    meta_path = out_dir / "meta.json"
    if meta_path.exists():
        try:
            old = json.loads(meta_path.read_text(encoding="utf-8"))
        except ValueError:
            old = {}
        if old.get("sha256") == h and old.get("release") == release:
            return False
    last = None
    for s in ds["series"]:
        e = _month_index(s["start"]) + len(s["values"]) - 1
        last = e if last is None or e > last else last
    meta = {
        "source": "World Bank Commodity Price Data (The Pink Sheet), monthly",
        "source_page": SOURCE_PAGE,
        "source_url": source_url,
        "release": release,
        "latest_month": _month_str(last) if last is not None else None,
        "fetched_at": fetched_at or _dt.datetime.now(_dt.timezone.utc).replace(microsecond=0).isoformat(),
        "series_count": len(ds["series"]),
        "sha256": h,
        "license": LICENSE,
    }
    if extra_meta:
        meta.update(extra_meta)
    (out_dir / "monthly.json").write_text(dumps_compact(ds) + "\n", encoding="utf-8")
    meta_path.write_text(json.dumps(meta, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return True


def collect(out_dir: pathlib.Path, page_url: str = SOURCE_PAGE, xlsx_path: pathlib.Path | None = None) -> bool:
    """Download (or read xlsx_path), parse, and write outputs. Returns True if data changed."""
    if xlsx_path is not None:
        p = pathlib.Path(xlsx_path)
        if p.stat().st_size > MAX_XLSX_BYTES:
            raise CollectError("too_big", "Workbook exceeds " + str(MAX_XLSX_BYTES) + " bytes.")
        url, data = "file:" + p.name, p.read_bytes()
    else:
        url, data = download(page_url)
    parsed = parse_workbook(data)
    ds = check_dataset(to_dataset(parsed["rows"]))
    return write_outputs(ds, out_dir, url, parsed["release"])


def check_dataset(ds: dict) -> dict:
    """Fail loudly unless the parsed data loads in the browser and has the core series."""
    from .analytics import AnalyticsError, validate_dataset

    try:
        validate_dataset(ds)
    except AnalyticsError as e:
        raise CollectError("bad_output", "Parsed data would not load in the dashboard: " + str(e)) from None
    ids = {s["id"] for s in ds["series"]}
    missing = [c for c in REQUIRED_SERIES if c not in ids]
    if missing:
        raise CollectError("unexpected_layout", "Workbook layout looks changed; missing series: " + ", ".join(missing)
                           + ". Check the header rows of the Monthly Prices sheet. Nothing was written.")
    return ds
