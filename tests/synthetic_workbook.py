"""Build small synthetic workbooks that mirror the layout of CMO-Historical-Data-Monthly.xlsx.

Values are made up and exist only to test the parser. They are not prices.
"""

from __future__ import annotations

import io

import openpyxl

HEADER_TITLE = [
    ["World Bank Commodity Price Data (The Pink Sheet)"],
    ["Updated on October 02, 2026"],
    ["Monthly prices in nominal US dollars, 1960 to present"],
    ["(monthly series are available only in excel)"],
]
NAMES = ["Crude oil, average", "Crude oil, Brent", "Crude oil, WTI", "Natural gas, US", "Natural gas, Europe",
         "Urea ", "Gold", "Silver", "Maize"]
UNITS = ["($/bbl)", "($/bbl)", "($/bbl)", "($/mmbtu)", "($/mmbtu)", "($/mt)", "($/troy oz)", "($/troy oz)", "($/mt)"]
CODES = ["CRUDE_PETRO", "CRUDE_BRENT", "CRUDE_WTI", "NGAS_US", "NGAS_EUR", "UREA_EE_BULK", "GOLD", "SILVER", "MAIZE"]


def synthetic_value(col: int, k: int):
    """Deterministic made-up value for column col at month offset k."""
    base = [2.0, 2.1, 1.9, 0.15, 0.4, 30.0, 35.0, 0.9, 45.0][col]
    return round(base * (1 + 0.01 * k) + 0.05 * ((k * (col + 3)) % 7), 4)


def build_rows(months: int = 30, with_codes: bool = True, start_year: int = 1960, gaps=True):
    rows = [list(r) for r in HEADER_TITLE]
    rows.append([None] + NAMES)
    rows.append([None] + UNITS)
    if with_codes:
        rows.append([None] + CODES)
    for k in range(months):
        y, m = start_year + k // 12, k % 12 + 1
        r = [f"{y}M{m:02d}"]
        for c in range(len(NAMES)):
            v = synthetic_value(c, k)
            if gaps and c == 4 and k < 5:
                v = ".."          # Europe gas starts later
            if gaps and c == 8 and k == 7:
                v = "…"      # missing month shown as an ellipsis
            r.append(v)
        rows.append(r)
    return rows


def build_xlsx(rows, sheet="Monthly Prices", extra_sheets=("Index", "Description")) -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = extra_sheets[0] if extra_sheets else sheet
    if extra_sheets:
        ws.append(["Index sheet"])
        for name in extra_sheets[1:]:
            wb.create_sheet(name).append(["notes"])
        ws = wb.create_sheet(sheet)
    for r in rows:
        ws.append(r)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
